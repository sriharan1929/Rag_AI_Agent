import os
from langchain_core.documents import Document

def _load_excel(path: str):
    """Load .xlsx / .xls — one LangChain Document per sheet."""
    ext  = os.path.splitext(path)[1].lower()
    docs = []

    if ext == ".xlsx":
        import openpyxl
        wb = openpyxl.load_workbook(path, data_only=True)
        for sheet_name in wb.sheetnames:
            ws   = wb[sheet_name]
            rows = []
            for row in ws.iter_rows(values_only=True):
                if any(cell is not None and str(cell).strip() for cell in row):
                    rows.append("\\t".join("" if c is None else str(c) for c in row))
            if rows:
                docs.append(Document(
                    page_content=f"Sheet: {sheet_name}\\n" + "\\n".join(rows),
                    metadata={"source": path, "sheet": sheet_name},
                ))

    elif ext == ".xls":
        import xlrd
        wb = xlrd.open_workbook(path)
        for sheet_name in wb.sheet_names():
            ws   = wb.sheet_by_name(sheet_name)
            rows = []
            for i in range(ws.nrows):
                row = ws.row_values(i)
                if any(str(c).strip() for c in row):
                    rows.append("\\t".join(str(c) for c in row))
            if rows:
                docs.append(Document(
                    page_content=f"Sheet: {sheet_name}\\n" + "\\n".join(rows),
                    metadata={"source": path, "sheet": sheet_name},
                ))

    return docs
